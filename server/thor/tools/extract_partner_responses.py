#!/usr/bin/env python3
import pandas as pd
import csv
import argparse

def extract_hyperlinks(excel_path):
    """
    Extract hyperlinks from Excel cells using openpyxl
    
    Returns:
        Dict mapping (sheet, row, col) -> hyperlink URL
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("  ⚠️  openpyxl not available - skipping link extraction")
        return {}
    
    try:
        wb = load_workbook(excel_path, data_only=False)
        links = {}
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.hyperlink:
                        key = (sheet_name, cell.row - 1, cell.column - 1)  # Convert to 0-indexed
                        links[key] = cell.hyperlink.target
        
        if links:
            print(f"  ✓ Extracted {len(links)} hyperlink(s) from Excel")
        
        return links
    except Exception as e:
        print(f"  ⚠️  Could not extract hyperlinks: {e}")
        return {}

def extract_responses(excel_path, output_path, control_list=None, application_type='SOFTWARE'):
    """
    Generic script to extract partner responses from Excel competency assessment files.
    
    Handles:
    - Multiple sheets with different column layouts
    - Single response controls (most sheets)
    - Multi-response controls (customer example sheets)
    - Automatic suffix mapping based on application_type (SERVICE or SOFTWARE)
    - Automatic detection of "ID" and "Partner Response" columns
    
    Args:
        excel_path: Path to Excel file
        output_path: Path to output CSV
        control_list: Optional list of controls to filter
        application_type: 'SERVICE' or 'SOFTWARE' for suffix mapping (default: SOFTWARE)
    """
    # Extract hyperlinks first
    print("  🔗 Extracting embedded hyperlinks...")
    hyperlinks = extract_hyperlinks(excel_path)
    
    excel_data = pd.read_excel(excel_path, sheet_name=None, header=None)
    responses = []
    
    # Controls that need suffix mapping for SERVICE/SOFTWARE variants
    suffix_controls = ['DOC-001', 'OPE-001', 'OPE-002', 'OPE-003', 
                      'PS-001', 'PS-002', 'PS-003', 'PS-004', 'PS-005',
                      'REL-001', 'REL-002', 'USE-CASE']
    
    # Create mapping for controls with suffix (always map for SERVICE/SOFTWARE)
    control_mapping = {}
    
    # Always create default mapping based on application_type
    for base_control in suffix_controls:
        control_mapping[base_control] = f"{base_control}-{application_type}"
    
    # Override with explicit control_list mappings if provided
    if control_list:
        for control in control_list:
            if control.endswith('-SERVICE') or control.endswith('-SOFTWARE'):
                base_control = control.replace('-SERVICE', '').replace('-SOFTWARE', '')
                control_mapping[base_control] = control
    
    for sheet_name, df in excel_data.items():
        if sheet_name.lower() == "introduction":
            continue
            
        # Find header row and ID column
        header_row = None
        id_col = None
        
        for idx, row in df.iterrows():
            for col_idx, cell in enumerate(row):
                if pd.notna(cell) and str(cell).strip() == "ID":
                    header_row = idx
                    id_col = col_idx
                    break
            if header_row is not None:
                break
        
        if header_row is None or id_col is None:
            continue
        
        # Define response columns based on sheet type
        response_cols = []
        
        # Check for known multi-response patterns first (application-type aware)
        if "GenAI Cust Ex Reqs" in sheet_name or "Generative AI Customer Example" in sheet_name:
            if application_type == 'SERVICE':
                response_cols = [3, 5, 7, 9]  # D, F, H, J (SERVICE)
            else:  # SOFTWARE
                response_cols = [5, 7, 9, 11]  # F, H, J, L (SOFTWARE)
        elif "Common Cust Example Reqs" in sheet_name or "Common Customer Example" in sheet_name:
            # SERVICE and SOFTWARE both use E, G, I, K for Common sheet
            response_cols = [4, 6, 8, 10]  # E, G, I, K (Partner Response for both SERVICE and SOFTWARE)
        else:
            # For single response sheets, find "Partner Response" column dynamically
            # Check row after header first
            if header_row + 1 < len(df):
                response_row = df.iloc[header_row + 1]
                for col_idx, cell in enumerate(response_row):
                    if pd.notna(cell) and "Partner Response" in str(cell):
                        response_cols.append(col_idx)
            
            # If not found, check the header row itself
            if not response_cols:
                header_row_data = df.iloc[header_row]
                for col_idx, cell in enumerate(header_row_data):
                    if pd.notna(cell) and "Partner Response" in str(cell):
                        response_cols.append(col_idx)
        
        if not response_cols:
            print(f"Skipping {sheet_name}: No Partner Response columns found")
            continue
        
        print(f"{sheet_name}: ID col={id_col}, Response cols={response_cols}")
            
        # Extract responses - start from different row based on sheet structure
        start_row = header_row + 2 if "Cust Ex Reqs" in sheet_name or "Customer Example" in sheet_name else header_row + 1
        
        for idx in range(start_row, len(df)):
            row = df.iloc[idx]
            control_id = row.iloc[id_col] if id_col < len(row) else None
            
            if pd.notna(control_id) and str(control_id).strip():
                control_id = str(control_id).strip()
                
                # Skip header-like entries
                if control_id and control_id not in ["ID", "Partner Response", "nan"]:
                    # Determine output control ID and inclusion
                    output_control_id = control_id
                    should_include = False
                    
                    # Apply suffix mapping for controls that need it
                    if control_id in suffix_controls:
                        output_control_id = f"{control_id}-{application_type}"
                    
                    if control_list is None:
                        should_include = True
                    else:
                        # Check exact match
                        if output_control_id in control_list:
                            should_include = True
                        # Check if base control matches
                        elif control_id in control_mapping:
                            should_include = True
                            output_control_id = control_mapping[control_id]
                    
                    if should_include:
                        # Extract from each response column
                        for resp_col in response_cols:
                            partner_response = row.iloc[resp_col] if resp_col < len(row) else None
                            response_text = str(partner_response).strip() if pd.notna(partner_response) else ""
                            
                            # Check for hyperlink in this cell
                            link_key = (sheet_name, idx, resp_col)
                            embedded_link = hyperlinks.get(link_key, "")
                            
                            if response_text and response_text not in ["nan", ""]:
                                responses.append({
                                    'controlId': output_control_id,
                                    'partner_response': response_text,
                                    'embedded_link': embedded_link
                                })
    
    # Save to CSV with hyperlink column
    with open(output_path, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=['controlId', 'partner_response', 'embedded_link'])
        writer.writeheader()
        writer.writerows(responses)
    
    print(f"Extracted {len(responses)} responses to {output_path}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='Extract partner responses from Excel to CSV')
    parser.add_argument('excel_path', help='Path to Excel file')
    parser.add_argument('output_path', help='Output CSV file path')
    parser.add_argument('--controls', nargs='*', help='List of control IDs to extract (optional)')
    parser.add_argument('--type', choices=['SERVICE', 'SOFTWARE'], default='SOFTWARE',
                       help='Application type: SERVICE or SOFTWARE (default: SOFTWARE)')
    
    args = parser.parse_args()
    extract_responses(args.excel_path, args.output_path, args.controls, args.type)
